#!/usr/bin/env python3
"""TOP_Code_Lookup.xlsx column D — the CCC 4-digit TOP code — agrees with the
program title beside it.

Column D is the bridge from MAP's integer TOP id to the COCI TOP codes that
kb/top_cip_map.json and the published crosswalk are keyed by, and the EACR's
CIP Sectors filter and row sections route through it (#1681). Nothing read the
column before 2026-09-24, and on 2026-09-25 it disagreed with its own column E
title on 81 of 198 rows: 35 codes that are not TOP codes at all (Fire
Technology at 2130, Drafting at 0909) put 611 cards in "No CIP assigned yet",
and 46 real-but-wrong codes (History at 2203, which is Ethnic Studies;
Nursing (RN) at 1211, which is Polysomnography) filed 328 cards under the
wrong sector. kb/_correct_top_lookup_code4.py corrected them; this test keeps
the column honest when someone next edits the workbook.

The rule is re-derived here from the TOP manual and the crosswalk rather than
imported from the correcting script, so a mistake in that script's rule cannot
pass its own check.

Imports the generator (for the CIP route), so it needs openpyxl. Run from the
repo root:

    python3 tests/top_code_lookup_code4_test.py
"""
import json
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from openpyxl import load_workbook  # noqa: E402

import excel_to_dashboard as gen  # noqa: E402

CHECKS = []


def check(name, ok, detail=""):
    CHECKS.append((name, bool(ok), detail))


# Rows whose column E title names a TOP code that column D deliberately does
# not carry, and why. Keep in step with KEEP in kb/_correct_top_lookup_code4.py.
TITLE_EXCEPTIONS = {
    "144": "title reads General Work Experience; the notes name college success skills (4930)",
}

# ─────────────────────────── the authorities ─────────────────────────────────
titles = {}  # lower-cased TOP title -> {4-digit category}
for code, v in json.load(open(os.path.join(ROOT, "kb", "reference", "top_categories.json"),
                              encoding="utf-8"))["codes"].items():
    titles.setdefault(v["title"].strip().lower(), set()).add(code[:4])
wb = load_workbook(os.path.join(ROOT, "kb", "reference", "topcip_2021_crosswalk.xlsx"),
                   read_only=True, data_only=True)
for row in wb.worksheets[0].iter_rows(values_only=True):
    if row and isinstance(row[0], (int, float)) and row[1]:
        titles.setdefault(str(row[1]).strip().lower(), set()).add(f"{float(row[0]):07.2f}"[:4])
wb.close()
real4 = set().union(*titles.values())

# ─────────────────────────── the lookup ──────────────────────────────────────
wb = load_workbook(os.path.join(ROOT, "TOP_Code_Lookup.xlsx"), data_only=True)
rows = {}
for r in wb["TOP Code Lookup"].iter_rows(min_row=2, max_col=7, values_only=True):
    if r[0] is not None:
        rows[str(r[0]).strip()] = {"code4": str(r[3]).strip() if r[3] else "",
                                   "title": str(r[4] or "").strip()}
wb.close()
check("the lookup loads its 198 MAP TOP ids", len(rows) >= 198, f"got {len(rows)}")

# 1. Every code is a TOP code.
blank = sorted(k for k, v in rows.items() if not v["code4"])
check("every row carries a column D code", not blank, f"blank: {blank}")
bogus = sorted((k, v["code4"]) for k, v in rows.items() if v["code4"] and v["code4"] not in real4)
check("every column D code is a TOP 4-digit category (manual or crosswalk)", not bogus,
      f"not TOP codes: {bogus[:12]}")

# 2. Where column E names a TOP title, column D is that title's category.
disagree = []
for k, v in rows.items():
    named = titles.get(v["title"].lower())
    if named and v["code4"] not in named and k not in TITLE_EXCEPTIONS:
        disagree.append((k, v["title"], v["code4"], sorted(named)))
check("column D agrees with the TOP title column E names", not disagree,
      "; ".join(f"{k} {t!r}: {c} but the title is {n}" for k, t, c, n in disagree[:10]))
stale = sorted(k for k in TITLE_EXCEPTIONS
               if k not in rows or rows[k]["code4"] in (titles.get(rows[k]["title"].lower()) or set()))
check("every title exception is still needed", not stale, f"retire: {stale}")

# 3. The worked failures stay fixed.
for mid, want, what in [("130", "2133", "Fire Technology (was 2130, no such code)"),
                        ("135", "2205", "History (was 2203, Ethnic Studies)"),
                        ("360", "2207", "Political Science (was 2205, History)"),
                        ("98", "1230", "Nursing (RN) (was 1211, Polysomnography)"),
                        ("66", "0956", "Welding Technology (was 0941, no such code)"),
                        ("46", "0953", "Drafting Technology (was 0909, no such code)")]:
    got = rows.get(mid, {}).get("code4")
    check(f"MAP TOP {mid} reads {want}: {what}", got == want, f"got {got!r}")

# 4. Every MAP TOP id reaches a CIP family through the generator's own route.
by_code4, _families = gen._load_cip_families()
unrouted = sorted(k for k, v in rows.items() if not by_code4.get(v["code4"]))
check("every MAP TOP id resolves to a CIP family", not unrouted, f"unresolved: {unrouted}")
lookup = gen._load_top_code_lookup()
check("History reaches CIP 54, not Ethnic Studies' 05",
      gen._cip_sector_for_tops(["135"], lookup, by_code4) == "54",
      f"got {gen._cip_sector_for_tops(['135'], lookup, by_code4)!r}")

# ─────────────────────────── report ──────────────────────────────────────────
passed = sum(1 for _, ok, _ in CHECKS if ok)
failed = [(n, d) for n, ok, d in CHECKS if not ok]
print(f"\ntop_code_lookup_code4_test — {passed}/{len(CHECKS)} checks passed")
for n, d in failed:
    print(f"  FAIL  {n}" + (f"\n        {d}" if d else ""))
sys.exit(1 if failed else 0)
