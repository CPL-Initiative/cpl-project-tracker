"""
IT, cybersecurity, data and AI credential catalog -> kb/it_ai_credential_catalog.json

Sam, 2026-09-30: "I am interested in all certs, but particularly emerging ones
from the tech giants ... Not just certs, but courses that result in badges or
certs as well." The catalog answers the input side of the CPL question for one
field: which industry credentials a student may walk in holding, who issues
each one, and whether any California community college has articulated it in
MAP yet. A credential no college has articulated is an opportunity, stated in
the terms a faculty member already uses.

Three sources, joined on (issuer, core name):

  kb/reference/cos_certifications.json
      CareerOneStop Certification Finder, the national list (USDOL ETA, data by
      Minnesota DEED), synced monthly by cos-authority-sync.yml. The broadest
      source and the one with the certifying organization on every row.
      Display requires the attribution in ATTRIBUTION below.
  kb/reference/credential_registry_national_sample.json
      Credential Engine's National Certification Collection, the 974-of-6,738
      hand capture Sam made on 2026-09-16. Adds issuers COS lacks.
  kb/reference/industry_credential_watch.json
      What the issuers themselves publish now: certifications, course
      certificates and badges, each verified against the issuer's own page on
      the date it records. The registries lag the issuers by a year or more on
      AI credentials; this file is what closes that gap, and it is the file the
      industry-credential watch agent keeps current.

  kb/reference/industry_credential_skills.json
      The skills and competencies each issuer publishes for a credential:
      Microsoft's "skills measured", AWS and CompTIA exam domains with their
      weights, the courses inside a Google Career Certificate. Faculty compare
      these against course outcomes. Harvested from the issuer's own exam guide,
      one exam version at a time; the watch agent extends coverage each run.

and one MAP side:

  credential_reference_data.js  (the CER, rebuilt daily)
      Every unified title with its issuer and its articulations, which carry
      the colleges and units. A catalog row joins to MAP when a unified title
      from the same issuer carries the same core name, or when ALIASES names
      the pairing outright.

Run from the repo root:
  python3 kb/_build_it_ai_credential_catalog.py            # write the JSON
  python3 kb/_build_it_ai_credential_catalog.py --xlsx PATH  # and a workbook

The output is a dated snapshot, not a guarded artifact: the CER it reads is
rebuilt every morning, so a committed catalog trails MAP by however many days
have passed since its `_inputs.cer_generated_at`. Rerun before citing a count.
"""
import argparse
import json
import os
import re
import sys
from collections import Counter, defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
COS = os.path.join(HERE, "reference", "cos_certifications.json")
CE = os.path.join(HERE, "reference", "credential_registry_national_sample.json")
WATCH = os.path.join(HERE, "reference", "industry_credential_watch.json")
SKILLS = os.path.join(HERE, "reference", "industry_credential_skills.json")
CIP_CSV = os.path.join(HERE, "reference", "CIPCode2020.csv")
CER_JS = os.path.join(ROOT, "credential_reference_data.js")
OUT = os.path.join(HERE, "it_ai_credential_catalog.json")

ATTRIBUTION = ("CareerOneStop (www.careeronestop.org), sponsored by the U.S. "
               "Department of Labor, Employment and Training Administration; "
               "data maintained by the Minnesota Department of Employment and "
               "Economic Development. Credential Engine Registry "
               "(credentialfinder.org), National Certification Collection.")

TECH_GIANTS = ["Google", "Microsoft", "Amazon Web Services", "IBM", "Oracle",
               "Cisco", "NVIDIA", "Salesforce", "Meta", "Apple"]

# (pattern on the source's organization string, canonical issuer, scope)
# scope "all": every credential the issuer offers is in the catalog.
# scope "it":  the issuer spans fields, so only credentials whose name places
#              them in an IT domain come in (Pearson's IT Specialist series,
#              not its Agriscience Foundations).
ISSUERS = [
    (r"^google|google cloud|grow with google", "Google", "all"),
    (r"microsoft", "Microsoft", "all"),
    (r"amazon", "Amazon Web Services", "all"),
    (r"^ibm\b|international business machines", "IBM", "all"),
    (r"^oracle|^mysql", "Oracle", "all"),
    (r"^cisco", "Cisco", "all"),
    (r"nvidia", "NVIDIA", "all"),
    (r"salesforce", "Salesforce", "all"),
    (r"^meta\b", "Meta", "all"),
    (r"^apple", "Apple", "all"),
    (r"comptia|computing technology industry", "CompTIA", "all"),
    (r"international information systems? security|\bisc2\b|\(isc\)", "ISC2", "all"),
    (r"information systems audit|\bisaca\b", "ISACA", "all"),
    (r"ec-council", "EC-Council", "all"),
    (r"global information assurance|\bgiac\b|sans giac", "GIAC", "all"),
    (r"linux foundation", "Linux Foundation", "all"),
    (r"linux professional institute", "Linux Professional Institute", "all"),
    (r"red hat", "Red Hat", "all"),
    (r"juniper", "Juniper Networks", "all"),
    (r"vmware", "VMware", "all"),
    (r"broadcom", "Broadcom", "all"),
    (r"citrix", "Citrix", "all"),
    (r"^f5$", "F5", "all"),
    (r"check point", "Check Point", "all"),
    (r"certified wireless network", "CWNP", "all"),
    (r"mile2", "Mile2", "all"),
    (r"certnexus", "CertNexus", "all"),
    (r"international association of privacy", "IAPP", "all"),
    (r"institute for operations research", "INFORMS", "all"),
    (r"sas institute", "SAS", "all"),
    (r"servicenow", "ServiceNow", "all"),
    (r"databricks", "Databricks", "all"),
    (r"snowflake", "Snowflake", "all"),
    (r"splunk", "Splunk", "all"),
    (r"python institute", "Python Institute", "all"),
    (r"institute for (the )?certification of computing", "ICCP", "all"),
    (r"cloud credential council", "Cloud Credential Council", "all"),
    (r"network appliance|netapp", "NetApp", "all"),
    (r"storage networking industry", "SNIA", "all"),
    (r"^suse", "SUSE", "all"),
    (r"teradata", "Teradata", "all"),
    (r"wireshark", "Wireshark", "all"),
    (r"certified internet web", "CIW", "all"),
    (r"ieee computer society", "IEEE Computer Society", "all"),
    (r"software testing qualifications", "ASTQB", "all"),
    (r"hewlett packard", "Hewlett Packard Enterprise", "all"),
    (r"^dell", "Dell Technologies", "all"),
    (r"^sap\b", "SAP", "all"),
    (r"^certiport", "Certiport", "it"),
    (r"^pearson", "Pearson", "it"),
    (r"international institute of business analysis", "IIBA", "it"),
    (r"eta international", "ETA International", "it"),
    (r"adobe", "Adobe", "it"),
    (r"project management institute|cognilytica", "PMI", "it"),
    (r"^opentext", "OpenText", "it"),
    (r"^codehs", "CodeHS", "it"),
    (r"isograd", "Isograd", "it"),
    (r"^treccert", "TRECCERT", "it"),
    (r"^sisa$", "SISA", "it"),
    (r"scaled agile", "Scaled Agile", "it"),
    (r"rocheston", "Rocheston", "it"),
    (r"disaster recovery institute", "DRI International", "it"),
    (r"smart automation certification alliance", "SACA", "it"),
]
_ISSUER_RX = [(re.compile(p, re.I), name, scope) for p, name, scope in ISSUERS]

# First match wins. AI before cyber ("AI Security Management" is an AI
# credential for this audience), cyber before data, business applications
# before cloud (an Oracle Fusion "Cloud Implementation Professional" is an ERP
# credential), data center before data.
DOMAINS = [
    ("AI and machine learning",
     r"artificial intelligence|machine learning|\bai\b|\bai-|generative|\bgenai|"
     r"watsonx|\bllms?\b|large language|prompt|agentforce|agentic|copilot|"
     r"vector search|deep learning|computer vision|\bml\b|mist ai|ai-native"),
    ("Cybersecurity",
     r"secur|cyber|hack|forensic|penetration|pentest|threat|incident|privacy|"
     r"\bsoc\b|defender|vulnerab|malware|encryption|intrusion|cissp|\bcism\b|"
     r"\bcisa\b|crisc|ccsp|sscp|cysa|resilience|cyberops|information assurance"),
    # "risk", "audit" and "governance" alone are not security words (PMI's Risk
    # Management Professional is a project credential); the security issuers'
    # risk and audit credentials arrive through SECURITY_ISSUERS instead.
    ("Business applications and productivity",
     r"implementation|functional consultant|dynamics 365|s/4hana|successfactors|"
     r"accounting|financials|\bhcm\b|procurement|supply chain|sales cloud|"
     r"service cloud|marketing|\berp\b|order management|payroll|human capital|"
     r"warehouse management|revenue management|inventory|office specialist|"
     r"\bexcel\b|\bword\b|outlook|powerpoint|\baccess\b|quickbooks|"
     r"power platform|power apps|power automate|workspace|\bic3\b|"
     r"digital literacy|key applications|living online|administrator\b.*salesforce|"
     r"salesforce certified|business analysis|product ownership|project management|"
     r"\bproject\+|scrum|agile"),
    ("Cloud, networks and systems",
     r"data center|datacenter"),
    ("Data and analytics",
     r"\bdata|analytic|database|\bsql\b|\bbi\b|business intelligence|power bi|"
     r"tableau|big data|warehouse|fabric|cognos|mysql|\bdb2\b|statistic|"
     r"predictive|snowpro|lakehouse|\bdatasys"),
    ("Cloud, networks and systems",
     r"cloud|azure|\baws\b|network|\bccna\b|\bccnp\b|\bccie\b|\bccde\b|\bccst\b|"
     r"\bcct\b|linux|unix|server|kubernetes|devops|virtual|vmware|storage|"
     r"wireless|routing|switching|administrator|system|infrastructure|"
     r"\baplus\b|a\+|it support|hardware|technician|endpoint|windows|devnet|"
     r"automation|hybrid|openshift|openstack|ansible|mainframe|z/os|"
     r"operating|\bdevice|messaging|teams|exchange|cloudnetx|meraki|citrix|"
     r"\bit fundamentals\b|\btech\+|\bitf\b|essentials"),
    ("Software development",
     r"develop|programm|java|python|javascript|html|css|c\+\+|software|\bweb\b|"
     r"coding|\bcode|\bapp\b|application|testing|tester|quality|computational|"
     r"blockchain|internet of things|\biot\b|unity|computer scien|\bux\b|"
     r"user experience"),
]
_DOMAIN_RX = [(name, re.compile(p, re.I)) for name, p in DOMAINS]
OTHER_IT = "Other IT"
# The domains a mixed-field issuer (scope "it") must land in to be catalogued.
IT_DOMAINS = {"AI and machine learning", "Cybersecurity", "Data and analytics",
              "Cloud, networks and systems", "Software development"}

# Issuers whose every credential is a security credential unless it is an AI
# one (a GIAC "Python Coder" is taught and tested as a security skill).
SECURITY_ISSUERS = {"GIAC", "ISC2", "EC-Council", "Mile2", "IAPP", "ISACA"}

# CIP sector (the two-digit CIP family) and the CIP 2020 code under it. Sam's
# ruling for the Exhibit Adoption view (2026-09-25) governs: "CIP sector is
# best ... We only need the CIP sector on this tab for filter and quick
# categorization ... I don't want to get sucked into the top/CIP black hole."
# So the sector is assigned here from the credential's domain and name, never
# voted from TOP, and it decides nothing about identity. First rule in the
# credential's domain whose pattern matches wins; the last rule of each domain
# is its default. Every code is checked against CIPCode2020.csv at build time.
CIP_RULES = {
    "AI and machine learning": [
        (r"secur|audit|governance|risk|privacy", "11.1003"),
        (r".", "11.0102"),
    ],
    "Cybersecurity": [
        (r"forensic|investigat", "43.0403"),
        (r".", "11.1003"),
    ],
    "Data and analytics": [
        (r"scien", "30.7001"),
        (r"database|\bsql\b|\bdba\b|data engineer|warehous|\bdb2\b|mysql|datasys|"
         r"lakehouse|snowpro|data integration|big data|cosmos|data platform", "11.0802"),
        (r"business (data )?analy", "30.7102"),
        (r".", "30.7101"),
    ],
    "Cloud, networks and systems": [
        (r"cloud|azure|\baws\b|devops|kubernetes|openshift|openstack|hybrid|"
         r"cloudnetx|solutions architect", "11.0902"),
        (r"network|\bccna\b|\bccnp\b|\bccie\b|\bccde\b|routing|switching|wireless|"
         r"meraki|junos|telecom|\b5g\b|data center", "11.0901"),
        (r"support|\ba\+|aplus|help desk|\btech\+|it fundamentals|\bitf\b|technician|"
         r"hardware|endpoint|device", "11.1006"),
        (r".", "11.1001"),
    ],
    "Software development": [
        (r"\bweb\b|html|css|javascript", "11.0801"),
        (r".", "11.0201"),
    ],
    "Business applications and productivity": [
        (r"office specialist|excel|\bword\b|outlook|powerpoint|\baccess\b|\bic3\b|"
         r"digital literacy|key applications|living online|publisher|workspace", "11.0601"),
        (r"project|scrum|agile", "11.1005"),
        (r".", "52.1201"),
    ],
    OTHER_IT: [(r".", "11.0103")],
}
_CIP_RX = {d: [(re.compile(p, re.I), c) for p, c in rules] for d, rules in CIP_RULES.items()}


def load_cip_titles(path):
    """code -> title for every live CIP 2020 code; a code the 2020 edition
    moved away or deleted is absent, so a rule naming one fails the build."""
    import csv
    titles = {}
    with open(path, encoding="utf-8-sig") as fh:
        for r in csv.DictReader(fh):
            if r["Action"].strip().lower() in ("moved from", "deleted"):
                continue
            code = r["CIPCode"].strip('="')
            titles[code] = r["CIPTitle"].strip().rstrip(".")
    return titles


def cip_of(name, domain):
    for rx, code in _CIP_RX[domain]:
        if rx.search(name):
            return code
    return None


def sector_title(t):
    small = {"and", "or", "of", "the", "in", "for"}
    words = t.lower().split()
    return " ".join(w if (i and w in small) else w[:1].upper() + w[1:]
                    for i, w in enumerate(words))


# MAP unified title -> the catalog name it is, where the names differ in more
# than punctuation, vendor prefix or an exam-code parenthetical. Each pairing
# was read by a person; never add one on a token overlap alone.
ALIASES = {
    ("Cisco", "Cisco Certified Network Associate (CCNA)"): "CCNA Certification",
    ("CompTIA", "CompTIA CySA+"): "CompTIA Cybersecurity Analyst",
    ("CompTIA", "CompTIA Cloud Essentials+"): "CompTIA Cloud Essentials",
    ("Amazon Web Services", "AWS Certified SysOps Administrator"):
        "AWS Certified SysOps Administrator - Associate",
    ("Linux Professional Institute", "LPIC-1 Linux Administrator"):
        "Linux Server Professional Certification - Level 1",
    ("Oracle", "Oracle Database SQL Certified Associate (1Z0-071)"):
        "Oracle Database SQL",
    ("Oracle", "Oracle Database PL/SQL Developer Certified Professional (1Z0-149)"):
        "Oracle Database Program with PL/SQL",
    ("Google", "Google IT Support Professional Certificate"):
        "Google IT Support Professional Certificate",
    # Microsoft Office Specialist editions (2016, 2019, 365) are one credential
    # for CPL purposes; MAP titles name the application, COS names the edition.
    ("Microsoft", "Microsoft Office Specialist — Excel"):
        "Microsoft Office Specialist: Excel Associate (Excel and Excel 2019)",
    ("Microsoft", "Microsoft Office Specialist — Word"):
        "Microsoft Office Specialist: Word Associate (Word and Word 2019)",
    ("Microsoft", "Microsoft Office Specialist — Outlook"):
        "Microsoft Office Specialist: Outlook Associate (Outlook and Outlook 2019)",
    ("Microsoft", "Microsoft Office Specialist — PowerPoint"):
        "Microsoft Office Specialist: PowerPoint Associate (PowerPoint and PowerPoint 2019)",
    ("Microsoft", "Microsoft Office Specialist — Excel Expert"):
        "Microsoft Office Specialist: Microsoft Excel Expert (Excel and Excel 2019)",
    ("Microsoft", "Microsoft Office Specialist — Word Expert"):
        "Microsoft Office Specialist: Microsoft Word Expert (Word and Word 2019)",
    ("Microsoft", "Microsoft Office Specialist — Access Expert"):
        "Microsoft Office Specialist: Microsoft Access Expert (Access and Access 2019)",
    ("Microsoft", "Microsoft Office Specialist — Access"):
        "Microsoft Office Specialist: Microsoft Access 2016",
}

# MAP titles that name a family rather than one credential. They stay out of
# the per-credential join and are counted apart, so a college's "CompTIA
# Certification (unspecified)" never lands on A+ by guesswork.
MAP_GENERIC = {"CompTIA Certification (unspecified)", "GIAC Certification",
               "Microsoft Office Specialist (generic)", "Microsoft Office Specialist"}

# Vendor prefixes and filler that differ between sources for one credential.
_PREFIX = re.compile(
    r"^(microsoft 365 certified|microsoft certified|microsoft|comptia|"
    r"aws certified|amazon web services|aws|google cloud certified|google cloud|"
    r"google|ibm certified|ibm|oracle certified|oracle|cisco certified|cisco|"
    r"red hat certified|red hat|nvidia certified|nvidia|salesforce certified|"
    r"salesforce|ec council|giac|isc2)\b[: -]*")


def norm(s):
    s = (s or "").lower()
    s = re.sub(r"[‐-―−]", "-", s)
    s = s.replace("&", " and ")
    s = re.sub(r"\([^)]*\)", " ", s)
    s = s.replace("+", "plus")
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return " ".join(s.split())


def core(name):
    """The name with vendor prefix, exam-code parentheticals and trailing
    'certification' stripped -- 'Microsoft Certified: Azure AI Fundamentals
    (AI-900)' and 'Microsoft Azure AI Fundamentals' share one core."""
    n = norm(name)
    for _ in range(2):
        n = _PREFIX.sub("", n).strip()
    n = re.sub(r"\b(certification|certificate|certified|exam)$", "", n).strip()
    return n


def issuer_of(org):
    for rx, name, scope in _ISSUER_RX:
        if rx.search(org or ""):
            return name, scope
    return None, None


def domain_of(name, issuer):
    text = name + " " + ("aws" if issuer == "Amazon Web Services" else "")
    for dname, rx in _DOMAIN_RX:
        if rx.search(text):
            if issuer in SECURITY_ISSUERS and dname != "AI and machine learning":
                return "Cybersecurity"
            return dname
    return "Cybersecurity" if issuer in SECURITY_ISSUERS else OTHER_IT


def load_json(path):
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def load_js_object(path):
    with open(path, encoding="utf-8") as fh:
        src = fh.read()
    return json.loads(src[src.index("{"):src.rindex("}") + 1])


def _add(rows, index, issuer, name, fields, source):
    key = (issuer, core(name))
    code = (fields.get("code") or "").upper().strip()
    hit = index.get(key) or (index.get((issuer, "#" + code)) if code else None)
    if hit is None:
        hit = {"issuer": issuer, "name": name, "sources": []}
        rows.append(hit)
    for k, v in fields.items():
        if v not in (None, "", []) and hit.get(k) in (None, "", []):
            hit[k] = v
    if source not in hit["sources"]:
        hit["sources"].append(source)
    index[key] = hit
    if code:
        index[(issuer, "#" + code)] = hit
    acr = norm(fields.get("acronym") or "")
    if acr:
        index.setdefault((issuer, "~" + acr), hit)
    return hit


def build(root=ROOT):
    cos = load_json(os.path.join(root, "kb", "reference", "cos_certifications.json"))
    ce = load_json(os.path.join(root, "kb", "reference",
                                "credential_registry_national_sample.json"))
    watch_path = os.path.join(root, "kb", "reference", "industry_credential_watch.json")
    watch = load_json(watch_path) if os.path.exists(watch_path) else {"credentials": []}
    skills_path = os.path.join(root, "kb", "reference", "industry_credential_skills.json")
    skills = load_json(skills_path) if os.path.exists(skills_path) else {"credentials": []}
    cip_titles = load_cip_titles(os.path.join(root, "kb", "reference", "CIPCode2020.csv"))
    for rules in CIP_RULES.values():
        for _p, code in rules:
            if code not in cip_titles:
                raise SystemExit(f"CIP_RULES names {code}, which is not a live CIP 2020 code")
    cer = load_js_object(os.path.join(root, "credential_reference_data.js"))

    rows, index = [], {}

    # The issuers' own pages first, so their current official names win.
    for w in watch.get("credentials", []):
        if not w.get("verified"):
            continue
        issuer, _ = issuer_of(w["issuer"])
        issuer = issuer or w["issuer"]
        hit = _add(rows, index, issuer, w["name"], {
            "code": w.get("code"), "kind": w.get("kind") or "certification",
            "level": w.get("level"), "status": w.get("status"),
            "launched": w.get("launched"), "emerging": bool(w.get("emerging")),
            "url": w.get("url"), "platform": w.get("platform"),
            "hours": w.get("hours"), "cost": w.get("cost"),
            "ace_credit": w.get("ace_credit"), "formerly": w.get("formerly"),
            "domain_hint": w.get("domain"), "notes": w.get("notes"),
            "verified_on": watch.get("_verified_on"),
        }, "issuer site")
        # A renamed credential keeps its old names as keys, so the registry
        # row and the MAP exhibit filed under the old name land on this row
        # (Cisco CyberOps Associate is CCNA Cybersecurity since 2026-02-03).
        for former in w.get("formerly") or []:
            index[(issuer, core(former))] = hit

    for c in cos["certifications"]:
        if not c.get("name"):
            continue
        issuer, scope = issuer_of(c.get("org"))
        if not issuer:
            continue
        if scope == "it" and domain_of(c["name"], issuer) not in IT_DOMAINS:
            continue
        _add(rows, index, issuer, c["name"], {
            "kind": "certification", "acronym": c.get("acronym"),
            "cos_type": c.get("type"), "url": c.get("url"), "cos_id": c.get("id"),
        }, "CareerOneStop")

    for c in ce["credentials"]:
        issuer, scope = issuer_of(c.get("provider"))
        if not issuer or c.get("provider_confidence") == "low":
            continue
        if scope == "it" and domain_of(c["name"], issuer) not in IT_DOMAINS:
            continue
        _add(rows, index, issuer, c["name"], {"kind": "certification"},
             "Credential Engine")

    # Domain: the watch file's own call where it made a specific one, else the
    # name. Its broad "IT" and "Productivity" calls defer to the name, which
    # tells a UX design certificate from an IT support one.
    hint_map = {"AI": "AI and machine learning", "Data": "Data and analytics",
                "Cyber": "Cybersecurity", "Cloud": "Cloud, networks and systems"}
    for r in rows:
        r["domain"] = hint_map.get(r.pop("domain_hint", None)) or domain_of(r["name"], r["issuer"])
        r["tech_giant"] = r["issuer"] in TECH_GIANTS
        r.setdefault("emerging", False)
        r.setdefault("kind", "certification")

    # ---- MAP side ---------------------------------------------------------
    catalog_issuers = {name for _p, name, _s in ISSUERS}
    map_generic, map_unmatched = [], []
    for u in cer["unified_titles"]:
        issuer, scope = issuer_of(u.get("issuer"))
        if issuer not in catalog_issuers:
            continue
        title = u["ut"]
        if scope == "it" and domain_of(title, issuer) not in IT_DOMAINS:
            continue
        if title in MAP_GENERIC:
            map_generic.append(title)
            continue
        colleges, units = set(), []
        for a in u.get("articulations", []):
            for loc in a.get("local", []):
                colleges.update(loc.get("colleges", []))
                if loc.get("u"):
                    units.append(float(loc["u"]))
        alias = ALIASES.get((issuer, title))
        hit = (index.get((issuer, core(alias))) if alias else None) \
            or index.get((issuer, core(title))) \
            or index.get((issuer, "~" + norm(re.sub(r"^.*?\(([^)]+)\)\s*$", r"\1", title))))
        if hit is None:
            # A MAP credential no registry lists (CompTIA Linux+, a retired
            # Microsoft exam, a Networking Academy course). It still belongs in
            # the catalog: colleges articulated it.
            hit = _add(rows, index, issuer, title, {"kind": "certification"}, "MAP")
            hit["domain"] = domain_of(title, issuer)
            hit["tech_giant"] = issuer in TECH_GIANTS
            hit["emerging"] = False
            map_unmatched.append(title)
        m = hit.setdefault("map", {"titles": [], "colleges": [], "units": [],
                                   "articulation_lines": 0})
        m["titles"].append(title)
        m["colleges"] = sorted(set(m["colleges"]) | colleges)
        m["units"] += units
        m["articulation_lines"] += u.get("n_articulation_lines") or 0
        if "MAP" not in hit["sources"]:
            hit["sources"].append("MAP")

    # Issuer-published skills: joined like everything else, on (issuer, exam
    # code) first and (issuer, core name) second.
    skills_attached = 0
    for sk in skills.get("credentials", []):
        if not sk.get("verified") or not sk.get("skills"):
            continue
        issuer, _ = issuer_of(sk["issuer"])
        issuer = issuer or sk["issuer"]
        code = (sk.get("code") or "").upper().strip()
        hit = (index.get((issuer, "#" + code)) if code else None) \
            or index.get((issuer, core(sk["name"])))
        if hit is None:
            continue
        hit["skills"] = {"kind": sk.get("skills_kind"), "exam_version": sk.get("exam_version"),
                         "domains": sk["skills"], "source_url": sk.get("source_url"),
                         "harvested_on": skills.get("_harvested_on")}
        skills_attached += 1

    for r in rows:
        code = cip_of(r["name"], r["domain"])
        r["cip"] = {"code": code, "title": cip_titles[code],
                    "sector": code[:2], "sector_title": sector_title(cip_titles[code[:2]]),
                    "source": "rule (domain and name)"}
        m = r.get("map")
        if m:
            u = m.pop("units")
            m["colleges_n"] = len(m["colleges"])
            m["units_min"] = min(u) if u else None
            m["units_max"] = max(u) if u else None
            r["map_status"] = "articulated" if m["colleges_n"] else "exhibit, not yet articulated"
        else:
            r["map_status"] = "not in MAP"

    rows.sort(key=lambda r: (r["domain"], not r["tech_giant"], r["issuer"],
                             -(r.get("map") or {}).get("colleges_n", 0), r["name"]))
    for i, r in enumerate(rows, 1):
        r["id"] = f"ITC{i:04d}"

    by_domain = defaultdict(Counter)
    for r in rows:
        by_domain[r["domain"]][r["map_status"]] += 1
    return {
        "_about": ("IT, cybersecurity, data and AI credentials -- certifications, "
                   "course certificates and badges -- by issuer, joined to MAP. "
                   "DERIVED by kb/_build_it_ai_credential_catalog.py; never edit "
                   "by hand. A 'not in MAP' row is an opportunity: a credential "
                   "students can hold that no college has an exhibit for."),
        "_attribution": ATTRIBUTION,
        "_inputs": {
            "cos_synced_at": cos.get("_synced_at"),
            "credential_engine_captured_on": ce.get("_captured_on"),
            "credential_engine_coverage": ce.get("_coverage"),
            "watch_verified_on": watch.get("_verified_on"),
            "cer_generated_at": cer.get("_generated_at"),
        },
        "_stats": {
            "credentials": len(rows),
            "issuers": len({r["issuer"] for r in rows}),
            "tech_giant_credentials": sum(r["tech_giant"] for r in rows),
            "emerging": sum(bool(r["emerging"]) for r in rows),
            "in_map": sum(r["map_status"] != "not in MAP" for r in rows),
            "articulated": sum(r["map_status"] == "articulated" for r in rows),
            "with_issuer_skills": skills_attached,
            "by_cip_sector": dict(Counter(f'{r["cip"]["sector"]} {r["cip"]["sector_title"]}'
                                          for r in rows).most_common()),
            "by_domain": {d: dict(c) for d, c in sorted(by_domain.items())},
            "map_titles_generic": sorted(map_generic),
            "map_titles_without_registry_row": sorted(map_unmatched),
        },
        "credentials": rows,
    }


def write_xlsx(cat, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill("solid", fgColor="002F6D")
    head_font = Font(bold=True, color="FFFFFF")
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    s = cat["_stats"]
    lines = [
        ("IT, cybersecurity, data and AI credentials by issuing agency", True),
        ("CPL Initiative, California Community Colleges Chancellor's Office", False),
        ("", False),
        (f"{s['credentials']:,} credentials from {s['issuers']} issuers. "
         f"{s['in_map']} are in MAP; colleges have articulated {s['articulated']}.", False),
        ("A row marked 'not in MAP' is a credential students can hold that no "
         "college has an exhibit for yet. Faculty decide the credit; the catalog "
         "names the evidence.", False),
        ("", False),
        ("Reading the columns", True),
        ("Domain and CIP sector group the credentials for filtering. The CIP sector is the "
         "two-digit CIP 2020 family and the CIP code the six-digit program code beneath it, "
         "both assigned from the credential's field and name; neither is an issuer's claim.", False),
        ("Emerging marks a credential launched or rebuilt since 2024, or one built for AI. "
         "Status comes from the issuer's own page: active, new, beta, retiring or retired. A "
         "retired credential stays listed because students still hold it.", False),
        ("MAP status reads 'articulated' when at least one college has articulated the "
         "credential, 'exhibit, not yet articulated' when MAP holds an exhibit no college has "
         "used, and 'not in MAP' otherwise.", False),
        ("Issuer-published skills are the exam domains, with the issuer's weights, or the "
         "courses inside a certificate, as the issuer states them.", False),
        ("", False),
        ("Sources", True),
        (f"CareerOneStop Certification Finder, synced {(cat['_inputs']['cos_synced_at'] or '')[:10]}", False),
        (f"Credential Engine National Certification Collection, sample captured "
         f"{cat['_inputs']['credential_engine_captured_on']}: 974 of its 6,738 members", False),
        (f"Issuer websites, checked {cat['_inputs']['watch_verified_on']}", False),
        (f"MAP platform articulated exhibits (CPL Credential Reference), "
         f"{(cat['_inputs']['cer_generated_at'] or '')[:10]}", False),
        ("", False),
        ("Attribution: " + cat["_attribution"], False),
        ("A snapshot. The live catalog rebuilds from the tracker repository "
         "(kb/_build_it_ai_credential_catalog.py).", False),
    ]
    for i, (t, bold) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(bold=bold, size=14 if i == 1 else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 110

    cols = [("Domain", 26), ("CIP sector", 30), ("CIP code", 40),
            ("Issuer", 22), ("Credential", 58), ("Exam or code", 13),
            ("Kind", 20), ("Level", 16), ("Tech giant", 10), ("Emerging", 10),
            ("Status", 10), ("MAP status", 26), ("Colleges articulating", 12),
            ("Units (range)", 13), ("MAP title", 44), ("Colleges", 60),
            ("Sources", 34), ("Official page", 48), ("Issuer-published skills", 90),
            ("Skills source", 48), ("Notes", 50)]

    def sheet(title, rows):
        w = wb.create_sheet(title)
        for j, (h, width) in enumerate(cols, 1):
            c = w.cell(row=1, column=j, value=h)
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(wrap_text=True, vertical="top")
            w.column_dimensions[get_column_letter(j)].width = width
        for i, r in enumerate(rows, 2):
            m = r.get("map") or {}
            units = ""
            if m.get("units_min") is not None:
                lo, hi = m["units_min"], m["units_max"]
                units = f"{lo:g}" if lo == hi else f"{lo:g}-{hi:g}"
            sk = r.get("skills") or {}
            skill_text = "; ".join(
                d["name"] + (f" ({d['weight']})" if d.get("weight") else "")
                for d in sk.get("domains", []))
            cip = r["cip"]
            vals = [r["domain"], f'{cip["sector"]} {cip["sector_title"]}',
                    f'{cip["code"]} {cip["title"]}',
                    r["issuer"], r["name"], r.get("code") or r.get("acronym") or "",
                    r.get("kind", ""), r.get("level") or r.get("cos_type") or "",
                    "yes" if r["tech_giant"] else "", "yes" if r["emerging"] else "",
                    r.get("status") or "", r["map_status"], m.get("colleges_n") or "",
                    units, "; ".join(m.get("titles", [])), "; ".join(m.get("colleges", [])),
                    ", ".join(r["sources"]), r.get("url") or "", skill_text,
                    sk.get("source_url") or "", r.get("notes") or ""]
            for j, v in enumerate(vals, 1):
                w.cell(row=i, column=j, value=v)
        w.freeze_panes = "F2"
        w.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
        return w

    rows = cat["credentials"]
    sheet("Tech giants, emerging",
          [r for r in rows if r["tech_giant"] and r["emerging"]])
    sheet("All credentials", rows)
    sheet("In MAP", sorted([r for r in rows if r.get("map")],
                           key=lambda r: -r["map"]["colleges_n"]))
    sheet("Not in MAP (opportunities)", [r for r in rows if not r.get("map")])

    w = wb.create_sheet("By issuer")
    hdr = ["Issuer", "Tech giant", "Credentials", "Emerging", "In MAP", "Articulated"]
    for j, h in enumerate(hdr, 1):
        c = w.cell(row=1, column=j, value=h)
        c.fill, c.font = head_fill, head_font
    agg = defaultdict(lambda: [0, 0, 0, 0])
    for r in rows:
        a = agg[r["issuer"]]
        a[0] += 1
        a[1] += bool(r["emerging"])
        a[2] += r["map_status"] != "not in MAP"
        a[3] += r["map_status"] == "articulated"
    for i, (iss, a) in enumerate(sorted(agg.items(), key=lambda kv: -kv[1][0]), 2):
        for j, v in enumerate([iss, "yes" if iss in TECH_GIANTS else ""] + a, 1):
            w.cell(row=i, column=j, value=v)
    w.column_dimensions["A"].width = 32
    w.freeze_panes = "A2"
    wb.save(path)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--xlsx", help="also write a workbook to this path")
    args = ap.parse_args(argv)
    cat = build()
    text = json.dumps(cat, indent=1, ensure_ascii=False) + "\n"
    with open(OUT, "w", encoding="utf-8") as fh:
        fh.write(text)
    s = cat["_stats"]
    print(f"wrote {os.path.relpath(OUT, ROOT)}: {s['credentials']} credentials, "
          f"{s['issuers']} issuers, {s['in_map']} in MAP, {s['articulated']} articulated")
    if args.xlsx:
        write_xlsx(cat, args.xlsx)
        print(f"wrote {args.xlsx}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
