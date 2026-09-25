#!/usr/bin/env python3
"""Correct TOP_Code_Lookup.xlsx column D (the CCC 4-digit TOP code) against
the TOP manual, and write a receipt of every row it changes.

WHY. Column D is the bridge from MAP's integer TOP id to the COCI TOP codes
every other reference file is keyed by. Nothing read it until the EACR's CIP
Sectors filter did (#1681, 2026-09-24), and measured on 2026-09-25 it
disagreed with its own column E program title on 81 of 198 rows: History at
2203 (Ethnic Studies), Political Science at 2205 (History), Fire Technology
at 2130 (no such code; 2133.00), Drafting at 0909 (0953.00), Nursing (RN) at
1211 (Polysomnography). 611 EACR cards sat in "No CIP assigned yet" because
their code was not a TOP code at all.

THE RULE. Column D is the 4-digit category of the TOP title that column E
names, read from the 7th Edition TOP Manual
(kb/reference/top_categories.json) and the published TOP-CIP crosswalk
(kb/reference/topcip_2021_crosswalk.xlsx, which carries the few titles the
manual extract lacks, such as 1119.00 Portuguese and 4932.00 General Work
Experience). Where column E names no TOP title exactly, HAND settles the row
from the manual and the row's own notes (column F), and says why. KEEP lists
the rows left as they stand, with the reason.

Rule 7 still governs what the code is for: it groups and filters, and never
determines identity.

    python3 kb/_correct_top_lookup_code4.py           # dry run: print the changes
    python3 kb/_correct_top_lookup_code4.py --apply   # rewrite column D + the receipt
"""
import json
import os
import sys

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LOOKUP = os.path.join(ROOT, "TOP_Code_Lookup.xlsx")
MANUAL = os.path.join(ROOT, "kb", "reference", "top_categories.json")
CROSSWALK = os.path.join(ROOT, "kb", "reference", "topcip_2021_crosswalk.xlsx")
RECEIPT = os.path.join(ROOT, "kb", "receipts", "top_code_lookup_code4_2026-09-25_s288.json")

# Rows whose column E names no TOP title exactly: the code, and why.
HAND = {
    "48": ("0708", "no exact TOP title; the notes name Cisco CCNA, so 0708.10 Computer Networking"),
    "54": ("0936", "no exact TOP title; the notes name printing, so 0936.00 Printing and Lithography"),
    "55": ("0952", "0958 is Water and Wastewater Technology; furniture and cabinet making is 0952.50 Mill and Cabinet Work"),
    "85": ("1107", "1108 is Japanese; Chinese is 1107.00"),
    "97": ("1230", "1220 is Speech/Language Pathology and Audiology; LVN is 1230.20 Licensed Vocational Nursing"),
    "98": ("1230", "1211 is Polysomnography; RN is 1230.10 Registered Nursing"),
    "101": ("1240", "1201 is Health Occupations, General; the notes name dental assisting, 1240.10 Dental Assistant"),
    "102": ("1240", "1250 is Emergency Medical Services; dental hygiene is 1240.20 Dental Hygienist"),
    "153": ("4932", "4935 is not a TOP code; 4932.00 General Work Experience"),
    "159": ("0112", "0106 is not a TOP code; 0112.00 Agriculture Business, Sales, and Service"),
    "160": ("0101", "0100 is not a TOP code; 0101.00 Agriculture Technology and Sciences, General"),
    "176": ("1117", "1130 is not a TOP code; Hmong falls under 1117.00 Asian, South Asian, and Pacific Islands"),
    "320": ("1222", "1290 is not a TOP code; 1222.00 Physical Therapist Assistant (the notes name OTA, 1218.00, in the same CIP family)"),
    "332": ("1199", "1134 is not a TOP code; Persian falls under 1199.00 Other Foreign Languages"),
    "406": ("3009", "2401 is not a TOP code; the notes name tourism, so 3009.00 Travel Services and Tourism"),
}

# Rows the rule would change but a curator should settle: the code stays.
KEEP = {
    "144": "column E reads General Work Experience (4932), but the notes name college success skills, "
           "which the existing 4930 fits; the title is the doubtful column",
}


def title_index():
    """Lower-cased TOP title -> set of 4-digit categories (manual, then crosswalk)."""
    idx = {}
    for code, v in json.load(open(MANUAL, encoding="utf-8"))["codes"].items():
        idx.setdefault(v["title"].strip().lower(), set()).add(code[:4])
    wb = load_workbook(CROSSWALK, read_only=True, data_only=True)
    for row in wb.worksheets[0].iter_rows(values_only=True):
        if row and isinstance(row[0], (int, float)) and row[1]:
            idx.setdefault(str(row[1]).strip().lower(), set()).add(f"{float(row[0]):07.2f}"[:4])
    wb.close()
    return idx


def plan():
    """[(row_number, map_top_id, title, notes, old, new, basis)] for every row that changes."""
    idx = title_index()
    wb = load_workbook(LOOKUP, data_only=True)
    ws = wb["TOP Code Lookup"]
    changes = []
    for n, row in enumerate(ws.iter_rows(min_row=2, max_col=7, values_only=True), start=2):
        if row[0] is None:
            continue
        mid = str(row[0]).strip()
        old = str(row[3]).strip() if row[3] else ""
        title = str(row[4] or "").strip()
        notes = str(row[5] or "").strip()
        if mid in KEEP:
            continue
        if mid in HAND:
            new, basis = HAND[mid]
        else:
            named = idx.get(title.lower())
            if not named:
                continue
            if len(named) != 1:
                raise SystemExit(f"row {mid}: title {title!r} names {sorted(named)}; settle it in HAND")
            new = next(iter(named))
            basis = f"column E names the TOP title {title!r}, category {new}"
        if new != old:
            changes.append((n, mid, title, notes, old, new, basis))
    wb.close()
    return changes


def main(argv):
    changes = plan()
    for _, mid, title, _, old, new, basis in changes:
        print(f"{mid:>4}  {old or '(blank)':>7} -> {new}  {title[:40]:40}  {basis[:70]}")
    print(f"{len(changes)} rows change")
    if "--apply" not in argv:
        return 0
    wb = load_workbook(LOOKUP)
    ws = wb["TOP Code Lookup"]
    for n, mid, _title, _notes, _old, new, _basis in changes:
        assert str(ws.cell(row=n, column=1).value).strip() == mid, (n, mid)
        ws.cell(row=n, column=4).value = new
    wb.save(LOOKUP)
    receipt = {
        "_about": "TOP_Code_Lookup.xlsx column D (CCC 4-digit TOP code) corrected against "
                  "kb/reference/top_categories.json and topcip_2021_crosswalk.xlsx by "
                  "kb/_correct_top_lookup_code4.py. Revert a row by writing `old` back to column D.",
        "session": "S288 SkyZ, 2026-09-25",
        "kept": KEEP,
        "rows": [dict(map_top_id=mid, title=title, notes=notes, old=old, new=new, basis=basis)
                 for _, mid, title, notes, old, new, basis in changes],
    }
    with open(RECEIPT, "w", encoding="utf-8") as f:
        json.dump(receipt, f, indent=1, ensure_ascii=False)
        f.write("\n")
    print(f"wrote {os.path.relpath(LOOKUP, ROOT)} and {os.path.relpath(RECEIPT, ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
