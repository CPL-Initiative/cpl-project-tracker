#!/usr/bin/env python3
"""Write the OpenClassrooms Digital Marketer crosswalk workbook from
kb/openclassrooms_out/crosswalk.json (run _build_openclassrooms_crosswalk.py first).

Ashley's ten columns lead the Crosswalk sheet in her order; the evidence columns a
faculty reviewer needs follow them.

Run:  python3 kb/_write_openclassrooms_workbook.py [OUT.xlsx]
"""
import datetime as dt
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "openclassrooms_out", "crosswalk.json")

NO_CR = "No MAP statewide credit recommendation exists"
NO_EX = "No existing MAP exhibit identified"
CAT_MAP = "1 - Existing MAP exhibit / articulation (local)"
CAT_ONMAP = "2 - Potential: course already on MAP through another exhibit"
CAT_NEW = "3 - Potential: COCI / catalog course not in MAP (new exhibit needs faculty review)"

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="002F6D")   # First Light seal blue
BAND = {"Direct match": "E3EDF9", "Strong": "EAF4EC", "Moderate": "FFFFFF", "Partial": "F5F5F5"}
thin = Side(style="thin", color="C8CED6")


def clean(s):
    s = "" if s is None else str(s)
    for _ in range(3):
        try:
            t = s.encode("cp1252").decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            break
        if t == s:
            break
        s = t
    return s.replace("_x000D_", " ").replace("\r", " ").strip()


def header(ws, cols, widths):
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 32


def body(ws, fills=None):
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(bottom=thin)
            if fills:
                f = fills.get(r[0].row)
                if f:
                    c.fill = PatternFill("solid", fgColor=f)


def main():
    d = json.load(open(SRC))
    today = dt.date.today()
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        HERE, "openclassrooms_out", f"{today:%Y%m%d}_OpenClassrooms_Digital_Marketer_CPL_Crosswalk.xlsx")
    wb = Workbook()

    # ------------------------------------------------------------ Crosswalk ---
    ws = wb.active
    ws.title = "Crosswalk"
    cols = ["Title", "Certificate Name", "Region", "Credit Recommendation Title", "Discipline",
            "Exhibit ID", "Exhibit Title", "College Name", "Course Name", "Course ID",
            "Opportunity Category", "Alignment Strength", "Topic Area", "OpenClassrooms Competencies Evidenced",
            "Units", "Credit Status", "Active Status", "How Found", "Course Description (COCI)",
            "COCI Control Number", "TOP Code", "Suggested Next Step"]
    widths = [22, 30, 18, 34, 22, 22, 30, 24, 32, 13, 30, 13, 24, 48, 7, 12, 22, 26, 70, 15, 26, 40]
    header(ws, cols, widths)
    fills = {}
    for m in d["map_rows"]:
        ws.append([d["title"], d["certificate"], m["region"],
                   f"{clean(m['credit_rec'])} (local credit recommendation; {NO_CR.lower()})",
                   clean(m["discipline"]), m["exhibit_id"], clean(m["exhibit_title"]), m["college"],
                   clean(m["course_title"]), clean(m["course_id"]), CAT_MAP, "Existing (related credential)",
                   "", "", m["units"], "", m["active"], f"MAP exhibit ({m['cpl_type']})", "",
                   m["ctrl"], clean(m["top"]),
                   "Faculty review: add OpenClassrooms Digital Marketer as an exhibit mapped to this "
                   "already-articulated course, or extend the existing exhibit."])
        fills[ws.max_row] = "FFF7E0"
    for c in d["courses"]:
        cat = CAT_ONMAP if c["on_map"] else CAT_NEW
        ex_title = NO_EX + (f" for OpenClassrooms (course is on MAP via: {clean(c['on_map'])})" if c["on_map"] else "")
        step = ("Faculty review of the OpenClassrooms portfolio and project deliverables against this "
                "course's SLOs; if approved, create a MAP exhibit.")
        if c["strength"] == "Direct match":
            step = ("Highest priority: confirm with Bakersfield College whether this series is the related "
                    "instruction for OpenClassrooms' apprenticeship; if so, create the MAP exhibit.")
        ws.append([d["title"], d["certificate"], c["region"], NO_CR, clean(c["discipline"]),
                   NO_EX, ex_title, c["college"], clean(c["course_title"]), clean(c["course_id"]),
                   cat, c["strength"], c["lens"], c["competencies"], c["units"], c["credit"],
                   c["active"], c["evidence"], clean(c["description"]), c["ctrl"], clean(c["top"]), step])
        fills[ws.max_row] = BAND[c["strength"]]
    body(ws, fills)
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    n_cross = ws.max_row

    # ------------------------------------------------------------- Programs ---
    wp = wb.create_sheet("Programs & Certificates")
    header(wp, ["College Name", "Region", "Program Title", "Award", "Program Fit", "TOP Code", "CIP Code",
                "COCI Status", "Units", "CTE"], [26, 18, 42, 40, 22, 11, 10, 12, 12, 6])
    order = {"Digital-marketing focused": 0, "Marketing / related": 1}
    for p in sorted(d["programs"], key=lambda p: (order[p["fit"]], p["region"], p["college"], p["program"])):
        wp.append([p["college"], p["region"], clean(p["program"]), p["award"], p["fit"], p["top"], p["cip"],
                   p["status"], p["units"], p["cte"]])
    body(wp)
    wp.auto_filter.ref = f"A1:J{wp.max_row}"

    # ------------------------------------------------------ MAP ACE recs ---
    wa = wb.create_sheet("MAP Related Credit Recs")
    header(wa, ["Credit Recommendation Title", "Exhibit ID", "Exhibit Title", "Source",
                "Colleges holding it in MAP", "Note"], [36, 16, 34, 22, 60, 50])
    for r in d["ace_recs"]:
        wa.append([r["credit_rec"], r["exhibit_id"], r["exhibit_title"], "ACE (military training)",
                   "; ".join(f"{c['college']} ({c['region']})" for c in r["colleges"]),
                   "Related subject only. No college course is articulated to it yet, and it credits "
                   "military training, so it does not transfer to OpenClassrooms. Useful as credit "
                   "recommendation wording faculty can reuse."])
    body(wa)

    # --------------------------------------------------- Competency map ---
    wc = wb.create_sheet("Competency Groups")
    header(wc, ["Competency Group (used in matching)", "Appendix A work processes"], [44, 30])
    for g in d["competencies"]:
        wc.append([g["group"], g["letters"]])
    body(wc)

    # -------------------------------------------------------------- Read Me ---
    wr = wb.create_sheet("Read Me", 0)
    wr.column_dimensions["A"].width = 46
    wr.column_dimensions["B"].width = 100
    rows = [
        ("OpenClassrooms Digital Marketer - California Community Colleges CPL Crosswalk", ""),
        ("Prepared", f"{today:%B %d, %Y} for Ashley (CPL Initiative, MAP team). A snapshot: COCI and MAP change daily."),
        ("Training reviewed", "OpenClassrooms Digital Marketer Registered Apprenticeship, O*NET 13-1161.01, RAPIDS 2077CB: "
                              "27 competencies (Appendix A), 400 hours of related instruction in seven online projects, "
                              "DOL Certificate of Completion of Apprenticeship."),
        ("", ""),
        ("Summary", "Count"),
        ("MAP statewide credit recommendations found", len(d["statewide_credit_recs"])),
        ("Existing MAP exhibit / articulation rows", f'=COUNTIF(Crosswalk!K2:K{n_cross},"1*")'),
        ("Potential: course already on MAP via another exhibit", f'=COUNTIF(Crosswalk!K2:K{n_cross},"2*")'),
        ("Potential: COCI course not in MAP", f'=COUNTIF(Crosswalk!K2:K{n_cross},"3*")'),
        ("  Direct match", f'=COUNTIF(Crosswalk!L2:L{n_cross},"Direct match")'),
        ("  Strong", f'=COUNTIF(Crosswalk!L2:L{n_cross},"Strong")'),
        ("  Moderate", f'=COUNTIF(Crosswalk!L2:L{n_cross},"Moderate")'),
        ("  Partial", f'=COUNTIF(Crosswalk!L2:L{n_cross},"Partial")'),
        ("Programs and certificates (Programs tab)", f"=COUNTA('Programs & Certificates'!A2:A{wp.max_row})"),
        ("", ""),
        ("Headline finding", "Bakersfield College offers APPR B73A-G, Digital Marketer Apprenticeship 1-7 (20 units, "
                             "credit, active in MIS Fall 2025). Each description restates one OpenClassrooms project "
                             "brief, in sequence. No MAP exhibit exists for it. Confirm with Bakersfield whether it is "
                             "OpenClassrooms' related instruction before any other outreach."),
        ("No statewide credit recommendation", "MAP holds no statewide (CCC Collaborative) credit recommendation for "
                                               "digital marketing, social media, SEO or advertising. Every existing MAP "
                                               "row here is a local exhibit for a different credential (CLEP, AMA, CFT, "
                                               "credit by exam) that is articulated to a related college course."),
        ("", ""),
        ("Opportunity categories", ""),
        (CAT_MAP, "The college already awards CPL into this course through MAP for a related credential. The "
                  "course is proven CPL-ready; OpenClassrooms would need its own exhibit or an extension."),
        (CAT_ONMAP, "The course aligns and is already on MAP for another credential, but not for OpenClassrooms."),
        (CAT_NEW, "Found in COCI (the statewide course inventory) with catalog descriptions. Not in MAP. A new "
                  "exhibit and faculty review would be required."),
        ("", ""),
        ("Alignment strength", ""),
        ("Direct match", "The course's description restates an OpenClassrooms project, course for course."),
        ("Strong", "The course is centered on what the apprenticeship teaches (digital, social, search, content, "
                   "email/CRM marketing, analytics) and its title plus catalog description evidence at least three "
                   "competency groups, or five from any marketing course."),
        ("Moderate", "Clear overlap on part of the training (principles of marketing, advertising, e-commerce, "
                     "market research), or a core-topic course whose description evidences fewer competencies."),
        ("Partial", "Overlap on one strand only (branding and public relations, UX design); worth a look where a "
                    "college has nothing stronger."),
        ("", ""),
        ("Method", "1) MAP: every exhibit and credit recommendation in MAP searched for marketing, advertising, "
                   "social media, SEO, e-commerce, UX, content and CRM. 2) COCI: all 141,738 COCI courses matched "
                   "on title AND a business, media or IT TOP code (TOP corroborates only; it never decides), plus "
                   "courses found by digital-marketing phrases in their catalog description. Internships, work "
                   "experience and independent study are excluded. 3) Each course's title and description are "
                   "compared against 13 competency groups built from the 27 Appendix A work processes (Competency "
                   "Groups tab). A title match alone never makes a row Strong."),
        ("Active status", "'Yes' means the course's control number appears in the MIS Fall 2025 course inventory. "
                          "Otherwise verify in the college's current catalog."),
        ("Discipline", "The Minimum Qualifications discipline where the subject code maps without ambiguity; "
                       "otherwise the TOP code, marked 'verify'."),
        ("Region", "Strong Workforce Program region (MAP's map_colleges roster)."),
        ("Sources", "COCI course list and program export (via the CPL Project Tracker), MIS Fall 2025 course "
                    "basic file, MAP exhibits and credit recommendations (read-only, 2026-09-29), the OpenClassrooms "
                    "Appendix A work process schedule and program syllabus supplied by Ashley."),
        ("Caveat", "This is a research list for faculty review, not an articulation decision. Catalog descriptions "
                   "are the college-entered COCI text; course outlines of record and SLOs should be read before a "
                   "faculty meeting."),
    ]
    for a, b in rows:
        wr.append([a, b])
    for r in wr.iter_rows():
        for c in r:
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wr["A1"].font = Font(name=FONT, size=14, bold=True, color="002F6D")
    for r in wr.iter_rows():
        if r[0].value in ("Summary", "Opportunity categories", "Alignment strength", "Headline finding"):
            r[0].font = Font(name=FONT, size=11, bold=True, color="002F6D")
            r[1].font = Font(name=FONT, size=10, bold=r[0].value == "Summary")
    for r in range(6, 15):
        wr.cell(row=r, column=2).alignment = Alignment(horizontal="left")

    # Excel computes the Read Me COUNTIFs on open (LibreOffice recalc stalls on this file here).
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    print(out, "crosswalk rows:", n_cross - 1)


if __name__ == "__main__":
    main()
